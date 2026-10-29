#!/usr/bin/env python3
"""Assemble complete slide images and speaker notes into a validated PowerPoint."""
import argparse
import hashlib
import json
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches
from PIL import Image


def assemble(images_dir: Path, content_file: Path, destination: Path) -> dict:
    content = json.loads(content_file.read_text())
    pages = content['pages']
    expected = list(range(1, len(pages) + 1))
    if [p['n'] for p in pages] != expected:
        raise ValueError('Content page numbers must be contiguous and ordered from 1')
    images = [images_dir / f'{n:02d}.png' for n in expected]
    for page, image in zip(pages, images):
        if not page.get('notes', '').strip():
            raise ValueError(f'Missing speaker notes on page {page["n"]}')
        if not image.is_file():
            raise FileNotFoundError(image)
        with Image.open(image) as im:
            if im.format != 'PNG' or abs(im.width/im.height - 16/9) > 0.02:
                raise ValueError(f'Expected a landscape 16:9 PNG: {image} {im.size}')
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    presentation.core_properties.title = content.get('title', 'Image presentation')
    presentation.core_properties.subject = 'Complete slide images with editable speaker notes'
    presentation.core_properties.keywords = 'image slides, speaker notes'
    for page, image in zip(pages, images):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        picture = slide.shapes.add_picture(str(image), 0, 0,
                                          width=presentation.slide_width,
                                          height=presentation.slide_height)
        picture.name = f'{page["n"]:02d} {page.get("question", "")}'
        picture._element.nvPicPr.cNvPr.set('descr', page.get('question', ''))
        slide.notes_slide.notes_text_frame.text = page['notes']
    destination.parent.mkdir(parents=True, exist_ok=True)
    temp = destination.with_name(destination.stem + '.building.pptx')
    presentation.save(temp)
    reread = Presentation(temp)
    if len(reread.slides) != len(pages):
        raise ValueError('Slide count changed during save')
    records = []
    for n, (slide, page, image) in enumerate(zip(reread.slides, pages, images), 1):
        if len(slide.shapes) != 1:
            raise ValueError(f'Page {n} should contain exactly one complete image')
        blob = slide.shapes[0].image.blob
        source_hash = hashlib.sha256(image.read_bytes()).hexdigest()
        if hashlib.sha256(blob).hexdigest() != source_hash:
            raise ValueError(f'Image changed in PowerPoint on page {n}')
        if slide.notes_slide.notes_text_frame.text != page['notes']:
            raise ValueError(f'Speaker notes changed on page {n}')
        records.append({'page': n, 'image_sha256': source_hash,
                        'notes_characters': len(page['notes'])})
    temp.replace(destination)
    return {'pptx': str(destination), 'slides': len(pages), 'speaker_notes': len(pages),
            'image_bytes_preserved': True, 'sha256': hashlib.sha256(destination.read_bytes()).hexdigest(),
            'pages': records}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('images_dir', type=Path)
    parser.add_argument('content_json', type=Path)
    parser.add_argument('output_pptx', type=Path)
    parser.add_argument('--report', type=Path, help='Optional verification JSON path')
    args = parser.parse_args()
    result = assemble(args.images_dir, args.content_json, args.output_pptx)
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(result, ensure_ascii=False, indent=2))
    print(json.dumps({k:v for k,v in result.items() if k != 'pages'}, ensure_ascii=False))

if __name__ == '__main__':
    main()
